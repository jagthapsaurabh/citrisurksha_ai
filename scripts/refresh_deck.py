"""Append the AI-upgrade slides (steps 2-12) to the board deck and place them
before the Thank-you slide. Idempotent: skips if the marker slide exists.

  pip install python-pptx
  python3 scripts/refresh_deck.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DECK = Path(__file__).resolve().parents[1] / "docs" / "pitch" / "CitriSuraksha_AI_Board_Pitch.pptx"

DEEP = RGBColor(0x07, 0x11, 0x08)
GREEN = RGBColor(0x11, 0x65, 0x30)
AMBER = RGBColor(0xFF, 0xB0, 0x00)
LIGHT = RGBColor(0xF4, 0xFB, 0xEF)
INK = RGBColor(0x1B, 0x27, 0x1C)
GRAY = RGBColor(0x5A, 0x6B, 0x58)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFB, 0xFF, 0xF8)
LINEC = RGBColor(0xDF, 0xEE, 0xDD)
MARKER = "MULTI-STAGE PIPELINE (AI v2)"


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.08):
    sh = s.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1.2)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = rad
    sh.shadow.inherit = False
    return sh


def txt(s, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, lines=None):
    tb = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Segoe UI"
    if lines:
        p.line_spacing = Pt(lines)
    return tb


def hdr(s, kicker, title, sub=None):
    rect(s, 0.63, 0.45, 0.52, 0.07, AMBER, shape=MSO_SHAPE.RECTANGLE)
    txt(s, 0.63, 0.56, 9, 0.3, kicker, 11, True, RGBColor(0xC7, 0x78, 0x00))
    txt(s, 0.6, 0.82, 12.2, 0.8, title, 26, True, GREEN)
    if sub:
        txt(s, 0.62, 1.5, 12.2, 0.5, sub, 12.5, False, GRAY)


def card(s, x, y, w, h, title, body):
    rect(s, x, y, w, h, CARD, line=LINEC)
    txt(s, x + 0.2, y + 0.14, w - 0.4, 0.5, title, 13.5, True, GREEN)
    txt(s, x + 0.2, y + 0.52, w - 0.4, h - 0.6, body, 10.8, False, GRAY, lines=14.5)


def footer(s, n):
    txt(s, 0.6, 7.08, 9, 0.3, "CitriSuraksha AI  ·  Scan · Detect · Protect  ·  Board Pitch — Confidential",
        9, False, RGBColor(0x9A, 0xAA, 0x97))
    txt(s, 12.3, 7.08, 0.5, 0.3, str(n), 9, True, RGBColor(0x9A, 0xAA, 0x97), align=PP_ALIGN.RIGHT)


def already_refreshed(prs) -> bool:
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and MARKER in sh.text_frame.text:
                return True
    return False


def build(prs):
    slides = []

    # ---- A: pipeline v2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    hdr(s, MARKER, "From one CNN to a six-signal decision engine",
        "Photo in, honest answer out - including 'I don't know'.")
    flow = ["OpenCV quality gate\nblur / light / corruption", "YOLO11 detector\nbox + crop (fallback localiser)",
            "Legacy CNN signal\n(bootstrap until trained)", "DINOv2 + Qdrant\nverified-image memory",
            "OpenCLIP zero-shot\nsecondary, always weak", "Decision engine\n6 honest outcomes"]
    x = 0.62
    for i, t in enumerate(flow):
        rect(s, x, 2.1, 1.92, 1.35, CARD, line=LINEC)
        txt(s, x + 0.08, 2.2, 1.76, 1.2, t, 10.5, i in (0, 5), GREEN if i in (0, 5) else INK, lines=14)
        if i < 5:
            a = s.shapes.add_shape(MSO_SHAPE.CHEVRON, In(x + 1.95), In(2.55), In(0.24), In(0.4))
            a.fill.solid(); a.fill.fore_color.rgb = AMBER; a.line.fill.background(); a.shadow.inherit = False
        x += 2.1
    card(s, 0.62, 3.8, 3.9, 1.5, "Decisions", "identified · uncertain · unknown · poor_image · no_pest_detected · non_target_image. Unknowns are stored for clustering.")
    card(s, 4.72, 3.8, 3.9, 1.5, "Cross-checks", "Model disagreement or a thin top-1/top-2 margin always downgrades to 'uncertain' - never a blind guess.")
    card(s, 8.82, 3.8, 3.9, 1.5, "Knowledge join", "AI identifies; PostgreSQL provides symptoms, prevention, cure, organic/chemical control in EN/MR/HI.")
    rect(s, 0.62, 5.6, 12.1, 0.95, GREEN)
    txt(s, 0.9, 5.75, 11.5, 0.7, "Saying 'I don't know' is a feature: it feeds expert review, which feeds the dataset, which feeds accuracy.",
        13.5, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slides.append(s)

    # ---- B: governance
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    hdr(s, "GOVERNANCE", "No unverified image trains. No unevaluated model deploys.")
    card(s, 0.62, 2.0, 3.9, 1.55, "Dataset versions", "Immutable manifests: class distribution, splits, dedupe report. Frozen test set can never leak into training.")
    card(s, 4.72, 2.0, 3.9, 1.55, "Frozen evaluation", "Accuracy, macro P/R/F1, per-class AP, confusion matrix, latency - on the frozen test only.")
    card(s, 8.82, 2.0, 3.9, 1.55, "Comparison gate", "Better overall but worse on one important pest = REJECT. Small regressions = manual review.")
    card(s, 0.62, 3.75, 3.9, 1.55, "Approve → Deploy", "Deploy requires frozen eval + comparison 'deploy'. Old production archived automatically.")
    card(s, 4.72, 3.75, 3.9, 1.55, "One-click rollback", "The previous production model is always redeployable in a single call.")
    card(s, 8.82, 3.75, 3.9, 1.55, "Bootstrap detector", "Verified classification images start YOLO11 box training (full-image bootstrap boxes, clearly labelled).")
    rect(s, 0.62, 5.6, 12.1, 0.95, DEEP)
    txt(s, 0.9, 5.75, 11.5, 0.7, "Statuses: training → testing → approved → production → rejected / archived. Every move is audited.",
        13, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slides.append(s)

    # ---- C: continuous improvement
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    hdr(s, "CONTINUOUS IMPROVEMENT", "The platform learns from every review")
    card(s, 0.62, 2.0, 5.95, 1.7, "Priority review queue (active learning)", "Unreviewed scans ranked by review_priority: low confidence, model disagreement, thin margins, visual novelty. Experts label what improves the model most.")
    card(s, 6.77, 2.0, 5.95, 1.7, "Novel-pest discovery", "Unidentified crops are DINOv2-embedded and clustered; coherent clusters surface as CANDIDATE NEW CLASS before anyone has named them.")
    card(s, 0.62, 3.9, 5.95, 1.7, "Calibrated priorities", "Review outcomes (AI right/wrong vs expert label) produce lift-based weight proposals for the queue - human-approved, never auto-applied.")
    card(s, 6.77, 3.9, 5.95, 1.7, "Farmer feedback loop", "Correct/wrong marks + admin corrections flow into verified labels; each dataset version retrains both classifier and detector.")
    rect(s, 0.62, 5.9, 12.1, 0.8, GREEN)
    txt(s, 0.9, 6.02, 11.5, 0.6, "More scans → smarter queue → better labels → sharper models → more trust → more scans.",
        13, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slides.append(s)

    # ---- D: validated & what we need
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    hdr(s, "VALIDATED & READY", "Steps 1-12 complete - what the board should know")
    card(s, 0.62, 2.0, 3.9, 1.55, "Latency in budget", "Per-stage budgets enforced; steady-state 0.004-0.04 s/image on CPU (report shipped).")
    card(s, 4.72, 2.0, 3.9, 1.55, "Security 7/7", "Upload caps, traversal blocked, auth enforced, no tracked secrets, leak-free diagnostics.")
    card(s, 8.82, 2.0, 3.9, 1.55, "65+ automated tests", "AI service + backend suites green; admin panel builds clean; release runbook documented.")
    card(s, 0.62, 3.75, 3.9, 1.55, "Zero paid AI", "Open-source PyTorch/TorchVision/DINOv2/OpenCLIP; all data and models on our servers.")
    card(s, 4.72, 3.75, 3.9, 1.55, "Ingest tooling ready", "Labelled-drive ingestion with dedupe + provenance turns collected images into trainable data.")
    card(s, 8.82, 3.75, 3.9, 1.55, "Open-source only", "No vendor lock-in; deployment on one VPS with or without Docker.")
    rect(s, 0.62, 5.6, 12.1, 1.15, DEEP)
    txt(s, 0.9, 5.72, 11.5, 0.95,
        "WE ASK: approve the labelled-image drive + 500-farmer Nagpur pilot; open doors to KVKs/departments.\nFirst trained models follow the runbook: dataset-v1 → frozen eval → compare → approve → deploy.",
        12.5, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, lines=18)
    slides.append(s)
    return slides


def main() -> int:
    prs = Presentation(str(DECK))
    if already_refreshed(prs):
        print("Deck already refreshed; nothing to do.")
        return 0
    new = build(prs)
    # rebuild slide order: [...old except thank-you] + new + [thank-you]
    xml_slides = prs.slides._sldIdLst
    ids = list(xml_slides)
    n = len(new)
    old, added = ids[:-n], ids[-n:]
    order = old[:-1] + added + [old[-1]]
    for el in list(xml_slides):
        xml_slides.remove(el)
    for el in order:
        xml_slides.append(el)
    # renumber footers in display order
    import re
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and re.fullmatch(r"\d+", sh.text_frame.text.strip()):
                sh.text_frame.paragraphs[0].runs[0].text = str(i)
    prs.save(str(DECK))
    print(f"Deck refreshed: {len(order)} slides.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
