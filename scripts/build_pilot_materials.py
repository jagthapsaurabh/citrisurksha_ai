"""Pilot prep pack (board-approved next moves):
  1. docs/pilot/Farmer_Onboarding_Sheet.docx  (EN/MR/HI one-pager)
  2. docs/pilot/Citrus_Field_Cards_MR_HI_EN.docx (printable trilingual pest cards)
  3. docs/pilot/KVK_Partnership_Deck.pptx (partnership variant of the board deck)

Capability-level language only: no internal model names, no paid/free/open-source
API framing (board confidentiality rules).

  pip install python-pptx python-docx
  python3 scripts/build_pilot_materials.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as In, Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "ai-service"))

from app.researcher_pest_guide import RESEARCHER_TRANSLATIONS as TR  # noqa: E402
from app.seed_20_pests import TWENTY_PESTS  # noqa: E402

OUT = ROOT / "docs" / "pilot"
OUT.mkdir(parents=True, exist_ok=True)

GREEN = RGBColor(0x11, 0x65, 0x30)
GRAY = RGBColor(0x5A, 0x6B, 0x58)

CARD_PESTS = ["citrus-psyllid", "citrus-leaf-miner", "citrus-blackfly", "citrus-whitefly",
              "brown-citrus-aphid", "citrus-mealybug", "citrus-red-mite", "fruit-fly"]
EN = {p["id"]: p for p in TWENTY_PESTS}


# ------------------------------------------------------------------ 1. onboarding
def build_onboarding():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Noto Sans"
    st.font.size = Pt(10.5)

    h = doc.add_heading("CitriSuraksha AI — Farmer Onboarding / शेतकरी नोंदणी / किसान पंजीकरण", level=1)
    for r in h.runs:
        r.font.color.rgb = GREEN

    doc.add_paragraph("Scan · Detect · Protect — citrus pest protection in your pocket.")
    p = doc.add_paragraph()
    p.add_run("स्कॅन · ओळख · संरक्षण — सिट्रस कीड संरक्षण तुमच्या खिशात. ").font.color.rgb = GRAY
    p.add_run("स्कैन · पहचान · सुरक्षा — सिट्रस कीट सुरक्षा आपकी जेब में.").font.color.rgb = GRAY

    doc.add_heading("What you get / तुम्हाला काय मिळेल / आपको क्या मिलेगा", level=2)
    for en, mr, hi in [
        ("Photo scan: point your camera at a pest or damaged leaf and get the pest name with treatment advice in seconds.",
         "फोटो स्कॅन: कीड किंवा झाडाला झालेली इजा दाखवा — कीडीचे नाव व उपचार काही सेकंदात.",
         "फोटो स्कैन: कीट या क्षतिग्रस्त पत्ती की फोटो लें — सेकंडों में कीट का नाम और उपचार।"),
        ("Prevention & cure: preventive and curative steps, organic and chemical options, safety notes.",
         "प्रतिबंध व उपचार: सेंद्रिय व रासायनिक पर्याय, सुरक्षा सूचना.",
         "रोकथाम व उपचार: जैविक व रासायनिक विकल्प, सुरक्षा सुझाव।"),
        ("Monthly Calendar of Operation, advisories, events and district pest alerts.",
         "मासिक कामकाज कॅलेंडर, सल्ले, कार्यक्रम व जिल्हा कीड सूचना.",
         "मासिक कैलेंडर, सलाह, कार्यक्रम व जिला कीट अलर्ट।"),
        ("Chat support: ask the assistant; an expert replies if needed.",
         "चॅट मदत: सहाय्यकाला विचारा; गरज असल्यास तज्ञ उत्तर देतात.",
         "चैट सहायता: सहायक से पूछें; ज़रूरत पर विशेषज्ञ जवाब देंगे।"),
    ]:
        doc.add_paragraph(en, style="List Bullet")
        p = doc.add_paragraph(mr, style="List Bullet")
        p.runs[0].font.color.rgb = GRAY
        p = doc.add_paragraph(hi, style="List Bullet")
        p.runs[0].font.color.rgb = GRAY

    doc.add_heading("Join in 4 steps / ४ स्टेप्स / 4 कदम", level=2)
    for en, mr, hi in [
        ("1. Install the app and open it. / अ‍ॅप इन्स्टॉल करा व उघडा. / ऐप इंस्टॉल करके खोलें।", "", ""),
        ("2. Register with your mobile number and farm details (village, crop, varieties).",
         "मोबाईल नंबर व शेती तपशील (गाव, पीक, वाण) नोंदवा.",
         "मोबाइल नंबर व खेती विवरण (गाँव, फसल, किस्म) दर्ज करें।"),
        ("3. Scan a leaf to see your first AI report. / पहिला AI अहवाल पाहण्यासाठी पान स्कॅन करा. / पहली AI रिपोर्ट देखने हेतु पत्ती स्कैन करें।", "", ""),
        ("4. Keep the app installed for pest alerts. / कीड सूचनांसाठी अ‍ॅप ठेवा. / कीट अलर्ट हेतु ऐप रखें।", "", ""),
    ]:
        line = en + (("  ·  " + mr + "  ·  " + hi) if mr else "")
        doc.add_paragraph(line, style="List Number")

    doc.add_heading("Good scan tips / चांगल्या स्कॅनसाठी टिप्स / अच्छे स्कैन टिप्स", level=2)
    for t in ["Daylight, steady hand. / उजेड, स्थिर हात. / दिन की रोशनी, स्थिर हाथ।",
              "Close-up of the insect or damaged part. / कीड किंवा इजेचा क्लोज-अप. / कीट या क्षतिग्रस्त हिस्से का क्लोज-अप।",
              "Fill the frame with the leaf/pest. / पान/कीड फ्रेमभर ठेवा. / पत्ती/कीट को फ्रेम में भरें।"]:
        doc.add_paragraph(t, style="List Bullet")

    p = doc.add_paragraph()
    p.add_run("Help / मदत / सहायता: ").bold = True
    p.add_run("Village extension worker / KVK helpline, or use Chat in the app. Language can be switched anytime in Profile (English / मराठी / हिंदी).")
    doc.save(str(OUT / "Farmer_Onboarding_Sheet.docx"))


# ------------------------------------------------------------------ 2. field cards
def build_field_cards():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Noto Sans"
    st.font.size = Pt(9.5)
    h = doc.add_heading("Citrus Field Cards — CitriSuraksha AI (EN · मराठी · हिंदी)", level=1)
    for r in h.runs:
        r.font.color.rgb = GREEN
    doc.add_paragraph("Print, cut and laminate. Verify chemical options with KVK / agriculture department before use.")

    for pid in CARD_PESTS:
        en = EN.get(pid, {})
        tr = TR.get(pid, {})
        mr, hi = tr.get("mr", {}), tr.get("hi", {})
        doc.add_heading(f"{en.get('common_name', pid)}  |  {mr.get('common_name', '')}  |  {hi.get('common_name', '')}", level=2)
        tbl = doc.add_table(rows=4, cols=2)
        tbl.style = "Light Grid Accent 1"
        rows = [
            ("Identify / ओळख / पहचान", f"EN: {en.get('symptoms', '')[:220]}\nMR: {mr.get('identification', '') or mr.get('symptoms', '')}\nHI: {hi.get('identification', '') or hi.get('symptoms', '')}"),
            ("Active period / काल / अवधि", f"MR: {mr.get('active_period', '-')}   HI: {hi.get('active_period', '-')}"),
            ("Prevent / प्रतिबंध / रोकथाम", f"MR: {mr.get('prevention', '')}\nHI: {hi.get('prevention', '')}"),
            ("If found / आढळल्यास / मिलने पर", f"MR: {mr.get('cure', '') or mr.get('biological_management', '')}\nHI: {hi.get('cure', '') or hi.get('biological_management', '')}\nETL: {mr.get('etl', '-')}"),
        ]
        for i, (k, v) in enumerate(rows):
            tbl.rows[i].cells[0].text = k
            tbl.rows[i].cells[1].text = v
            for para in tbl.rows[i].cells[0].paragraphs:
                for run in para.runs:
                    run.bold = True
        doc.add_paragraph("")
    doc.save(str(OUT / "Citrus_Field_Cards_MR_HI_EN.docx"))


# ------------------------------------------------------------------ 3. KVK deck
PGREEN = PRGB(0x11, 0x65, 0x30)
PAMBER = PRGB(0xFF, 0xB0, 0x00)
PLIGHT = PRGB(0xF4, 0xFB, 0xEF)
PDEEP = PRGB(0x07, 0x11, 0x08)
PGRAY = PRGB(0x5A, 0x6B, 0x58)
PWHITE = PRGB(0xFF, 0xFF, 0xFF)
PCARD = PRGB(0xFB, 0xFF, 0xF8)
PLINE = PRGB(0xDF, 0xEE, 0xDD)


def krect(s, x, y, w, h, fill=None, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, In(x), In(y), In(w), In(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = PPt(1.2)
    sh.adjustments[0] = 0.08
    sh.shadow.inherit = False
    return sh


def ktxt(s, x, y, w, h, text, size=14, bold=False, color=PGRAY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lines=None):
    tb = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = PPt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Segoe UI"
    if lines:
        p.line_spacing = PPt(lines)
    return tb


def khdr(s, kicker, title, sub=None):
    krect(s, 0.63, 0.45, 0.52, 0.07, fill=PAMBER)
    ktxt(s, 0.63, 0.56, 9, 0.3, kicker, 11, True, PRGB(0xC7, 0x78, 0x00))
    ktxt(s, 0.6, 0.82, 12.2, 0.8, title, 27, True, PGREEN)
    if sub:
        ktxt(s, 0.62, 1.52, 12.2, 0.5, sub, 12.5, False, PGRAY)


def kcard(s, x, y, w, h, title, body):
    krect(s, x, y, w, h, fill=PCARD, line=PLINE)
    ktxt(s, x + 0.2, y + 0.14, w - 0.4, 0.5, title, 13.5, True, PGREEN)
    ktxt(s, x + 0.2, y + 0.5, w - 0.4, h - 0.55, body, 10.8, False, PGRAY, lines=14.5)


def build_kvk_deck():
    prs = Presentation()
    prs.slide_width = In(13.333)
    prs.slide_height = In(7.5)
    B = prs.slide_layouts[6]

    def base(fill=PLIGHT):
        s = prs.slides.add_slide(B)
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False
        return s

    s = base(PDEEP)
    ktxt(s, 0.9, 2.2, 11.5, 1.2, "CitriSuraksha AI × Krishi Vigyan Kendra", 44, True, PWHITE)
    ktxt(s, 0.9, 3.5, 11.5, 0.6, "Partnership proposal — citrus protection at district scale", 20, True, PRGB(0xD9, 0xF9, 0xDF))
    ktxt(s, 0.9, 4.6, 11.5, 0.8, "Scan · Detect · Protect  ·  Nagpur, Maharashtra", 14, False, PRGB(0xC9, 0xD9, 0xC6))

    s = base()
    khdr(s, "WHY TOGETHER", "Extension reach × always-on field intelligence")
    kcard(s, 0.62, 2.1, 3.9, 2.3, "The gap", "One extension officer serves thousands of orchards; a wrong or late diagnosis spreads faster than advice.")
    kcard(s, 4.72, 2.1, 3.9, 2.3, "The platform", "Photo diagnosis in seconds, trilingual advisory, month-wise care calendar, district alerts, expert chat.")
    kcard(s, 8.82, 2.1, 3.9, 2.3, "The partner", "KVK scientists validate, advise and turn field data into research and training outcomes.")

    s = base()
    khdr(s, "PILOT MODEL", "500 farmers · Nagpur district · 6 months")
    kcard(s, 0.62, 2.1, 5.95, 1.6, "KVK role", "Nodal scientist; verification & labelling camps; advisory review; ETL validation on field cards.")
    kcard(s, 6.77, 2.1, 5.95, 1.6, "Our role", "App rollout & training, AI operations, support desk, alert campaigns, reporting.")
    kcard(s, 0.62, 3.9, 5.95, 1.6, "Joint", "Monthly review of pest incidence dashboard; co-branded advisories; camp calendar.")
    kcard(s, 6.77, 3.9, 5.95, 1.6, "Farmers", "Free to use; consent-based data; language of their choice; priority expert support.")

    s = base()
    khdr(s, "WHAT KVK GAINS", "Tools that amplify, not replace, the scientist")
    kcard(s, 0.62, 2.1, 3.9, 2.3, "Live surveillance", "Geo-tagged, time-stamped pest incidence from hundreds of orchards - outbreaks visible weekly, not seasonally.")
    kcard(s, 4.72, 2.1, 3.9, 2.3, "Research-grade data", "Expert-verified, stage-labelled image dataset with provenance - material for papers, theses and trials.")
    kcard(s, 8.82, 2.1, 3.9, 2.3, "Extension leverage", "AI triages the routine; scientists spend time on the cases that need them. Alerts carry KVK guidance.")

    s = base()
    khdr(s, "GOVERNANCE & TRUST", "Every label verified; every model evaluated")
    kcard(s, 0.62, 2.1, 3.9, 2.3, "Verification first", "Farmer images never train the AI until an expert verifies them. Unknowns are flagged, not guessed.")
    kcard(s, 4.72, 2.1, 3.9, 2.3, "Evaluation gates", "New models pass frozen-test evaluation and side-by-side comparison before deployment; rollback is one click.")
    kcard(s, 8.82, 2.1, 3.9, 2.3, "Data residency", "All images, labels and models stay on our servers; farmer consent; co-ownership terms in the MoU.")

    s = base()
    khdr(s, "SIX-MONTH PLAN", "Camps, data, first trained models")
    kcard(s, 0.62, 2.1, 3.9, 2.3, "Month 0-1", "MoU; nodal scientist; 500-farmer onboarding via camps; field-card printing; app training.")
    kcard(s, 4.72, 2.1, 3.9, 2.3, "Month 1-3", "Labelled-image drive (20 pests + negative); weekly verification; first dataset version; alert channel live.")
    kcard(s, 8.82, 2.1, 3.9, 2.3, "Month 3-6", "First evaluated models deployed; district alert pilot; joint review + public results brief.")

    s = base(PDEEP)
    ktxt(s, 0.9, 2.4, 11.5, 1.0, "We request", 30, True, PAMBER)
    ktxt(s, 0.9, 3.3, 11.5, 2.2,
         "1. A nodal scientist and two labelling-camp dates.\n2. Advisory review slot for the first trilingual advisories.\n3. MoU covering data co-ownership and farmer consent.",
         16, False, PWHITE, lines=26)
    ktxt(s, 0.9, 5.6, 11.5, 0.6, "CitriSuraksha AI · Scan · Detect · Protect", 13, True, PRGB(0x84, 0xBD, 0x00))

    prs.save(str(OUT / "KVK_Partnership_Deck.pptx"))


def main() -> int:
    build_onboarding()
    build_field_cards()
    build_kvk_deck()
    print("pilot materials written to", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
