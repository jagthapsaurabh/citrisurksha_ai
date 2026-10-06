"""Board deck copy polish: human voice, no version numbers or implementation
counts, no ops-ready/Docker line, no em/en dashes.

  pip install python-pptx
  python3 scripts/humanize_deck.py
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

DECK = Path(__file__).resolve().parents[1] / "docs" / "pitch" / "CitriSuraksha_AI_Board_Pitch.pptx"

# key = distinctive opening of the paragraph (first chars), value = full new text
REWRITES = {
    "An AI-powered citrus pest detection": "A phone app that tells a farmer which pest is on the tree and what to do about it, with departments and researchers reading the same data.",
    "Born in Nagpur": "Born in Nagpur, the Orange City of India 🍊",
    "Pests and diseases are the single largest": "Pests and diseases cause most of the preventable loss in citrus, and the first two days matter most.",
    "20–40% crop losses": "Up to 40% crop losses",
    "FAO estimates that plant pests": "The FAO says pests and diseases destroy up to 40% of food crops worldwide every year. Citrus alone faces more than 20 major pests.",
    "Extension officers serve thousands": "One extension officer serves thousands of farmers. A wrong or late diagnosis spreads from village to village before anyone can check it.",
    "Mis-identification leads to wrong molecules": "When the pest is misidentified, farmers spray the wrong chemical. Money is wasted, residue builds up and friendly insects die.",
    "Asian citrus psyllid vectors": "The psyllid that spreads greening disease can only be caught early by trained eyes in the field, and there are never enough of them.",
    "Scientific advisories rarely reach": "Good advisories rarely reach farmers in Marathi or Hindi at the right week of the crop.",
    "Departments and universities lack": "Departments and universities have no live picture of where pests are appearing, so action comes after the outbreak.",
    "Source: FAO": "Source: FAO.",
    "Board Pitch — Confidential": "Board Pitch, Confidential",
    "A complete digital protection layer": "One platform that connects the farmer's phone to the department's desk.",
    "Android & iOS (Expo)": "Android and iOS. The farmer scans a leaf and gets the pest name with severity, symptoms and prevention plus cure advice, in English, Marathi or Hindi. History, care calendar, advisories, chat and alerts are built in.",
    "FastAPI + PostgreSQL + in-house": "A Python backend with its own vision service. The classifier covers the 20 citrus pests plus a 'no pest' class and improves with every verified upload. Push alerts reach a whole district at once.",
    "Role-based web panel (admin, agronomist": "A role based web panel for admins, agronomists, labelers and support staff: verify or correct AI results, publish advisories, manage pests and insecticides, train the models, send campaigns and answer farmer chats.",
    "Scan → Detect → Protect in under 30": "From scan to advice in under 30 seconds, with a human expert always behind it.",
    "React Native (Expo SDK 54)": "React Native\nAndroid + iOS\nOffline-safe, 3 languages",
    "FastAPI + SQLAlchemy": "Python backend\nSQL database\nSecure logins, fast responses",
    "Admin SDK push": "Push alerts\ndevice tokens\ndelivery logs",
    "Admin & Expert Panel — React": "️ Admin & Expert Panel, a React web app hosted independently",
    "The vision model is trained in-house": "The vision model is trained in house on our own verified images, so the data and the models belong to us.",
    "21-class CNN returns": "The classifier returns the most likely pests with confidence and severity",
    "New model version registered": "A new model version goes live and accuracy climbs every cycle",
    "The loop never stops": "The loop keeps turning: more scans, more verified images, a sharper model, more trust, more scans.",
    "Simple login (phone + password)": "Simple login with phone and password, and a dashboard where everything is one tap away.",
    "Onboarding in minutes": "Onboarding takes minutes: phone login and a farm profile with village, crop and varieties.\nOne dashboard with every service: Scan, History, Pest Guide, Calendar, Blog, insecticides, Events and Chat.\nAlerts appear on the bell the moment they arrive.",
    "Camera or gallery": "Camera or gallery: the photo is uploaded and answered in seconds.\nA complete report with pest name, severity, stage, symptoms and preventive plus curative control.\nEvery report carries a safety note on registered pesticides.",
    "Knowledge compiled from a 33-page": "Knowledge taken from a 33 page scientific citrus pest compendium and advisories written by researchers.",
    "Pest Guide: all 20 citrus pests": "Pest Guide: all 20 citrus pests with symptoms, prevention, cure and organic plus chemical control.\nCalendar of Operation: month wise orchard care, managed by the admin team.\nBlog and advisories published straight from the admin panel.\nCIB and RC recommended insecticide reference for each pest.",
    "Chatbot first, human expert second": "Chatbot first, human expert second, and district wide push alerts when it matters.",
    "Real-time chat": "Real time chat: the bot answers at once, unresolved queries reach the admin panel and agronomists reply with attachments.\nPush campaigns deliver district pest alerts and advisories even when the app is closed.\nAn in-app inbox keeps alerts with read and unread status.",
    "One tap — English": "One tap: English, मराठी, हिंदी",
    "The full experience (UI, pest names": "The whole experience, from menus to pest names and cure advice, is localised per farmer and switchable anytime.",
    "Trilingual from day one": "Trilingual from day one: English, Marathi and Hindi ship in the app and pest knowledge carries researcher checked translations.\nThe farmer's own profile: village, district, land, varieties and irrigation, so advice can follow the crop stage.\nLanguage choice is stored on the server, so advisories and blogs arrive in the farmer's language.\nAdding another Indian language later is a content job, not a rebuild.",
    "One dashboard to run the entire ecosystem": "One dashboard to run the whole system. Numbers shown are demo data.",
    "Live KPIs: users, pests, detections": "Live KPIs for users, pests, detections, the unreviewed queue, training images, feedback and alert-ready farmers.\nRole based menus: admin, agronomist, labeler and support each see only their own workspace.\nEvery KPI card opens the screen where the work happens.",
    "Every farmer scan lands here": "Every farmer scan lands here for verification. This is where the AI's accuracy is actually made.",
    "Agronomists see the farmer's photo": "The agronomist sees the farmer's photo next to the AI's best guesses.\nOne click to verify, correct the pest or stage, and attach a scientist approved note.\nApprove and Train turns a reviewed upload into a labelled training image on the spot.\nFarmers see Verified or Corrected status in their history, so trust is visible.",
    "Model training as an operational routine": "Model training as a routine operation, not a research project.",
    "Strict 21-class discipline": "Strict discipline: no production model is registered until verified images exist for all 20 pests plus the negative class.\nA live coverage board shows exactly which pest still needs photos.\nModel profiles can be chosen per training run and custom pests can be added.\nResearcher and browser data feeds keep enriching the knowledge base.",
    "Compose once": "Compose once and every registered farmer device receives the alert.\nFirebase status and delivery logs are built into the panel for accountability.\nUsed for outbreak alerts, weather advisories, events and new advisories.\nDelivery numbers feed back into engagement reporting.",
    "Transfer-training": "Training runs from the panel or a command line, with jobs and metrics.",
    "Strict mode: model registry refuses": "The registry refuses to promote a model until every class has verified images.",
    "Model versions are registered": "Model versions are listed and switchable from the AI Models screen.",
    "Backend normalises uploads": "Uploads are resized before inference and every call is timed and guarded.",
    "Before enough real images exist": "Before enough real images exist, a cautious bootstrap answers plausible photos and rejects blank ones, so the app is never dead.",
    "Competitors can copy a model": "Anyone can copy a model.\nNobody can copy a labelled,\nlocation tagged, expert verified\ncitrus pest dataset grown\nseason after season with\nfarmers and departments.\n\nThat dataset is the moat.",
    "From \"what is eating my crop?\"": "From 'what is eating my crop?' to a named pest with severity, without travel or waiting.",
    "CIB-RC aligned, stage-wise": "CIB and RC aligned chemical plus organic options, matched to the stage, cut wasted spray spend.",
    "Month-wise Calendar of Operation and alerts": "A month wise Calendar of Operation and timely alerts stop outbreaks before damage.",
    "Chatbot + real agronomist": "Chatbot plus real agronomist replies in the farmer's own language.",
    "Every scan saved": "Every scan is saved, building a crop health history that helps with loans and insurance.",
    "Bottom line for the farmer": "Bottom line for the farmer: fewer losses, lower input cost, safer produce, and the expert now comes to them.",
    "Live, geo-tagged pest incidence": "Live, location tagged pest incidence from thousands of scans; outbreak alerts go to a whole district in one click; extension staff triage with AI pre screening.",
    "A growing, expert-verified": "A growing, expert verified, stage labelled image dataset plus researcher data feeds: raw material for papers, trials and varietal studies.",
    "Evidence-based demand signals": "Evidence based demand signals: which pests, where and when, enabling responsible product guidance aligned with CIB and RC.",
    "Time-stamped, photo-verified": "Time stamped, photo verified crop damage history reduces risk in assessment and claims.",
    "Season-over-season pest maps": "Season over season pest maps inform spray programs, nursery policy and greening containment.",
    "India is among the world's leading": "India is among the world's leading citrus producers, and Vidarbha is home to the famous Nagpur orange. Citrus is high value and high risk, exactly where careful detection pays.",
    "Rural smartphone penetration": "Rural smartphones and cheap data mean we need no hardware rollout. The distribution channel already sits in the farmer's pocket.",
    "Digital agriculture, AI for farming": "Digital agriculture and plant surveillance are national priorities, and departments are actively looking for technology partners.",
    "The pipeline (scan": "The routine of scan, verify, train and advise works for any crop. Citrus is where we start; other horticulture crops can follow.",
    "This is not a slide deck": "This is not a slide deck. The platform is built.",
    "A production-grade MVP": "A working product across app, backend and admin panel, ready for a pilot.",
    "Farmer app — Expo SDK 54": "Farmer app for Android and iOS in three languages, ready to ship.\nBackend with role based auth, review workflows, live chat and push logging.\nAI service with its own training pipeline and model registry.\nAdmin panel covering every operation from review to campaigns.",
    "Full IP & data ownership": "All IP and data owned by us; every model trained in house.\nA cautious bootstrap keeps the app useful before the first training run.\nA human expert is in the loop on every detection.\nSecurity basics done: token auth, hashed passwords, secrets kept out of code.",
    "App + backend + admin + AI pipeline built": "App, backend, admin and AI pipeline built; knowledge digitised.",
    "PHASE 1 · 0–6 MO": "PHASE 1 · FIRST 6 MONTHS",
    "PHASE 2 · 6–12 MO": "PHASE 2 · MONTH 6 TO 12",
    "PHASE 3 · 12–24 MO": "PHASE 3 · MONTH 12 TO 24",
    "500-farmer pilot": "A pilot with 500 farmers and a KVK or university partner; a labelled image drive for all 21 classes; the first trained models.",
    "Top-1 accuracy ≥ 90%": "At least 90% top accuracy on field images; a light on device model; automatic outbreak clustering from location data.",
    "State-wide rollout": "State wide rollout in Maharashtra, more citrus regions, a second crop playbook and data access for agencies.",
    "KPIs we will report": "We will report to the board every quarter: model accuracy, verified image coverage, active farmers, scans per week, alert delivery rate and estimated loss avoided.",
    "Budget for field image collection": "Budget for field image collection, labelling camps with university students and agronomist review time. That is the fuel of the flywheel.",
    "• A defensible, data-backed": "• A defensible, data backed AI asset in Indian horticulture\n• A public good platform with strong CSR and policy alignment\n• Recurring value: agency dashboards, certified advisories, research datasets\n• First mover brand: the app that protects oranges",
    "Photo in, honest answer out": "Photo in, honest answer out, including 'I don't know'.",
    "identified · uncertain · unknown": "Identified, uncertain, unknown, poor image, no pest, or not a citrus photo. Unknowns are stored for clustering.",
    "Model disagreement or a thin": "If models disagree, or the top two guesses are close, the answer downgrades to 'uncertain'. Never a blind guess.",
    "AI identifies; PostgreSQL provides": "The AI identifies the pest; the database provides symptoms, prevention, cure and control advice in three languages.",
    "Saying 'I don't know' is a feature": "Saying 'I don't know' is a feature. It feeds expert review, which feeds the dataset, which feeds accuracy.",
    "Accuracy, macro P/R/F1": "Accuracy, precision, recall, F1, per class scores, confusion matrix and latency, measured on the frozen test only.",
    "Better overall but worse": "Better overall but worse on one important pest means reject. Small regressions mean manual review.",
    "Deploy requires frozen eval": "Deployment requires a frozen test evaluation and a positive comparison. The old production model is archived automatically.",
    "Statuses: training": "Statuses move from training to testing, approved, production, and finally rejected or archived. Every move is audited.",
    "Unreviewed scans ranked": "Unreviewed scans are ranked: low confidence, model disagreement, close guesses, unusual images. Experts label what improves the model most.",
    "Unidentified crops are embedded": "Unidentified crops are grouped by visual similarity; a coherent cluster surfaces as a candidate new class before anyone has named it.",
    "Review outcomes (AI right/wrong": "Review outcomes, where the AI was right or wrong against the expert label, produce proposals for queue weights. A human approves them, never the machine.",
    "Correct/wrong marks + admin corrections": "Correct and wrong marks plus admin corrections flow into verified labels; each dataset version retrains both the classifier and the detector.",
    "More scans → smarter queue": "More scans, a smarter queue, better labels, sharper models, more trust, more scans.",
    "Steps 1-12 complete": "Everything on the previous pages is built and tested",
    "Per-stage budgets enforced": "Every stage has a time budget and stays within it on a plain CPU server.",
    "Upload caps, traversal blocked": "Upload limits, blocked path traversal, enforced logins, no secrets in code, clean diagnostics.",
    "65+ automated tests": "Automated tests",
    "AI service + backend suites green": "AI and backend test suites pass, the admin panel builds clean and the release runbook is written down.",
    "Labelled-drive ingestion": "Ingestion with dedupe and provenance turns collected images into trainable data.",
    "No vendor lock-in": "Everything runs on servers we control.",
    "WE ASK: approve the labelled-image drive": "WE ASK: approve the labelled image drive and the 500 farmer Nagpur pilot, and open doors to KVKs and departments.\nFirst trained models follow the runbook: build a dataset, evaluate on the frozen test, compare, approve, deploy.",
    # second-pass anchors (text already rewritten once; clear stale bullets below)
    "The vision model is trained in house": "The vision model is trained in house on our own verified images, so the data and the models belong to us.",
    "In-house vision engine, trained only": "In-house vision engine, trained only on our own verified dataset.\nThe classifier covers the 20 citrus pests plus a no pest class and returns its best guesses with confidence.",
    "Farmer app for Android and iOS": "Farmer app for Android and iOS in three languages, ready to ship.\nBackend with role based auth, review workflows, live chat and push logging.\nAI service with its own training pipeline and model registry.\nAdmin panel covering every operation from review to campaigns.",
    "All IP and data owned by us": "All IP and data owned by us; every model trained in house.\nA cautious bootstrap keeps the app useful before the first training run.\nA human expert is in the loop on every detection.\nSecurity basics done: token auth, hashed passwords, secrets kept out of code.",
    "Onboarding takes minutes": "Onboarding takes minutes: phone login and a farm profile with village, crop and varieties.\nOne dashboard with every service: Scan, History, Pest Guide, Calendar, Blog, insecticides, Events and Chat.\nAlerts appear on the bell the moment they arrive.",
    "Camera or gallery:": "Camera or gallery: the photo is uploaded and answered in seconds.\nA complete report with pest name, severity, stage, symptoms and preventive plus curative control.\nEvery report carries a safety note on registered pesticides.",
    "Pest Guide: all 20 citrus pests": "Pest Guide: all 20 citrus pests with symptoms, prevention, cure and organic plus chemical control.\nCalendar of Operation: month wise orchard care, managed by the admin team.\nBlog and advisories published straight from the admin panel.\nCIB and RC recommended insecticide reference for each pest.",
    "Real time chat:": "Real time chat: the bot answers at once, unresolved queries reach the admin panel and agronomists reply with attachments.\nPush campaigns deliver district pest alerts and advisories even when the app is closed.\nAn in-app inbox keeps alerts with read and unread status.",
    "Trilingual from day one:": "Trilingual from day one: English, Marathi and Hindi ship in the app and pest knowledge carries researcher checked translations.\nThe farmer's own profile: village, district, land, varieties and irrigation, so advice can follow the crop stage.\nLanguage choice is stored on the server, so advisories and blogs arrive in the farmer's language.\nAdding another Indian language later is a content job, not a rebuild.",
    "Live KPIs for users, pests": "Live KPIs for users, pests, detections, the unreviewed queue, training images, feedback and alert-ready farmers.\nRole based menus: admin, agronomist, labeler and support each see only their own workspace.\nEvery KPI card opens the screen where the work happens.",
    "The agronomist sees the farmer's photo": "The agronomist sees the farmer's photo next to the AI's best guesses.\nOne click to verify, correct the pest or stage, and attach a scientist approved note.\nApprove and Train turns a reviewed upload into a labelled training image on the spot.\nFarmers see Verified or Corrected status in their history, so trust is visible.",
    "Strict discipline:": "Strict discipline: no production model is registered until verified images exist for all 20 pests plus the negative class.\nA live coverage board shows exactly which pest still needs photos.\nModel profiles can be chosen per training run and custom pests can be added.\nResearcher and browser data feeds keep enriching the knowledge base.",
    "Compose once and every registered": "Compose once and every registered farmer device receives the alert.\nFirebase status and delivery logs are built into the panel for accountability.\nUsed for outbreak alerts, weather advisories, events and new advisories.\nDelivery numbers feed back into engagement reporting.",
    "From one CNN": "From a single classifier to a six signal decision engine",
    "Legacy CNN signal": "Established classifier signal\n(careful bootstrap until trained)",
    "Ops-ready": "",
}

GLOBAL = [
    (re.compile(r"\s—\s"), ", "),
    (re.compile(r"^—\s*"), ""),
    (re.compile(r"\s—$"), ""),
    (re.compile(r"\s–\s"), " to "),
    (re.compile(r"(?i)expo\s*sdk\s*\d+"), ""),
    (re.compile(r"\b\d+\.\d+\.\d+(?:\.\d+)?\b"), ""),
    (re.compile(r"(?i)(fastapi|sqlalchemy|torchvision|pytorch)\s+\d[\d.]*"), r"\1"),
    (re.compile(r"\b12 API modules\b"), "API modules"),
    (re.compile(r"\b15 modules\b"), "modules"),
    (re.compile(r"\b18 screens\b"), "screens"),
    (re.compile(r"\s{2,}"), "  "),
]


def set_para(p, text: str) -> None:
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run().text = text


def main() -> int:
    prs = Presentation(str(DECK))
    changed = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            paras = sh.text_frame.paragraphs
            for idx, p in enumerate(paras):
                text = "".join(r.text for r in p.runs)
                if not text.strip():
                    continue
                new = None
                for key, val in REWRITES.items():
                    if text.strip().startswith(key):
                        new = val
                        break
                if new is None:
                    new = text
                    for rx, rep in GLOBAL:
                        new = rx.sub(rep, new)
                    if new != text:
                        changed += 1
                        set_para(p, new)
                else:
                    # whole-box rewrite: first paragraph gets the new text and
                    # stale bullet paragraphs below it are cleared
                    changed += 1
                    set_para(p, new)
                    for q in paras[idx + 1:]:
                        if "".join(r.text for r in q.runs).strip():
                            set_para(q, "")
                    break
    prs.save(str(DECK))
    print("paragraphs rewritten:", changed)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
